"""
Preenchimento automático de planilha de presença a partir de um sistema web.

Para cada dia do mês, o script abre a agenda do sistema, lê a tabela de
"Realizados/Cancelados", identifica os pacientes com atendimento realizado e
marca "X" na planilha (um "X" por atendimento; "XX" se houve dois no dia).

Veja o README.md para instruções completas.
"""

import calendar
import os
import unicodedata

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

# ======================= CONFIGURAÇÃO =======================
ANO, MES = 2026, 10
ARQUIVO = "pacientes.xlsx"     # use uma CÓPIA da planilha nos testes
COL_NOME = 1                   # coluna A = nomes
COL_DIA_1 = 2                  # coluna B = dia 1, C = dia 2, ...
ESPECIALIDADE = "NOME DA ESPECIALIDADE"   # exatamente como aparece no <select> do site

DIA_INICIO = 1
DIA_FIM = None                 # None = mês todo. Para teste, use 2

# Endereço do sistema. Pode ser definido aqui ou pela variável de ambiente SISTEMA_URL
URL = os.environ.get("SISTEMA_URL", "https://SEU-SISTEMA.exemplo.com.br/")
# ============================================================


def normalizar(texto):
    """Minúsculas, sem acentos e sem espaços repetidos (para comparar nomes)."""
    texto = unicodedata.normalize("NFD", texto or "")
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return " ".join(texto.lower().split())


def eh_atendido(status):
    s = normalizar(status)
    return ("atendimento realizado" in s) or ("at_atendido" in s)


def achar_linha(ws, nome):
    """Devolve a linha do paciente na planilha; cria uma nova se não existir."""
    alvo = normalizar(nome)
    for row in range(2, ws.max_row + 1):
        valor = ws.cell(row=row, column=COL_NOME).value
        if valor and normalizar(str(valor)) == alvo:
            return row
    nova = ws.max_row + 1
    ws.cell(row=nova, column=COL_NOME, value=nome.strip())
    return nova


JS_LER_TABELA = """
() => {
  const norm = s => (s || '').replace(/\\s+/g, ' ').trim();
  const resultado = [];
  const tabelas = [...document.querySelectorAll('table')]
    .filter(t => !t.querySelector('table'));
  for (const t of tabelas) {
    const trs = [...t.querySelectorAll('tr')];
    const cab = trs.find(tr => {
      const x = norm(tr.innerText);
      return x.includes('Paciente') && x.includes('Status') && x.includes('Evolu');
    });
    if (!cab) continue;
    const titulos = [...cab.children].map(c => norm(c.innerText));
    const iPac = titulos.findIndex(x => x === 'Paciente');
    const iSta = titulos.findIndex(x => x === 'Status');
    if (iPac < 0 || iSta < 0) continue;
    for (const tr of trs.slice(trs.indexOf(cab) + 1)) {
      const cel = [...tr.children];
      if (cel.length < titulos.length) continue;
      const nome = norm(cel[iPac].innerText);
      const c = cel[iSta];
      const partes = [c.innerText, c.textContent];
      c.querySelectorAll('*').forEach(e => {
        partes.push(e.value, e.alt, e.title, e.getAttribute('aria-label'),
                    e.getAttribute('src'), e.innerText);
      });
      const status = norm(partes.filter(Boolean).join(' '));
      if (nome) resultado.push({nome: nome, status: status});
    }
  }
  return resultado;
}
"""


def escolher_especialidade(page):
    """Seleciona a especialidade avisando o site (evento 'change')."""
    resultado = page.evaluate("""(nome) => {
      const sels = [...document.querySelectorAll('select')];
      const sel = sels.find(s => [...s.options].some(o => o.text.trim() === nome));
      if (!sel) return 'NENHUM select com essa opcao. Selects na pagina: ' +
        sels.map(s => (s.id || s.name || '?') + '[' +
          [...s.options].map(o => o.text.trim()).join('|') + ']').join(' ; ');
      const op = [...sel.options].find(o => o.text.trim() === nome);
      sel.value = op.value;
      sel.dispatchEvent(new Event('input', {bubbles: true}));
      sel.dispatchEvent(new Event('change', {bubbles: true}));
      if (window.jQuery) window.jQuery(sel).val(op.value).trigger('change');
      return 'ok id=' + sel.id + ' valor=' + op.value;
    }""", ESPECIALIDADE)
    if not resultado.startswith("ok"):
        print("Especialidade:", resultado)
    page.wait_for_timeout(800)


def abrir_tela(page):
    """Menu: Atendimento > Horário Marcado."""
    page.get_by_role("link", name="paciente Atendimento").click()
    page.get_by_role("link", name="Horário Marcado").click()
    page.wait_for_selector("#btnBuscar")
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(2000)
    escolher_especialidade(page)


def buscar_dia(page, dia):
    data = f"{dia:02d}/{MES:02d}/{ANO}"

    page.get_by_role("img", name="...").click()                   # abre o calendário
    page.get_by_role("link", name=str(dia), exact=True).click()   # escolhe o dia
    escolher_especialidade(page)                                  # reforça a especialidade
    page.get_by_role("img", name="buscar").click()                # lupa

    page.wait_for_timeout(1500)
    if page.get_by_text("Não foi possível executar a instrução SQL").count() > 0:
        raise Exception("o site devolveu erro de SQL")

    try:
        page.wait_for_selector(f"text=Realizados/Cancelados do dia {data}", timeout=8000)
    except Exception:
        return None   # título não apareceu: tratado no main

    page.wait_for_load_state()
    page.wait_for_timeout(1000)
    return page.evaluate(JS_LER_TABELA)


def main():
    wb = load_workbook(ARQUIVO)
    ws = wb.active
    ultimo = DIA_FIM or calendar.monthrange(ANO, MES)[1]
    falhas = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page = browser.new_page()
        page.goto(URL)
        input("Faça o LOGIN no navegador e, quando entrar no sistema, "
              "pressione Enter aqui...")

        abrir_tela(page)

        for dia in range(DIA_INICIO, ultimo + 1):
            registros = None
            erro = None

            # até 2 tentativas por dia
            for tentativa in (1, 2):
                try:
                    registros = buscar_dia(page, dia)
                    erro = None
                except Exception as e:
                    registros = None
                    erro = e
                if registros is not None:
                    break
                try:
                    abrir_tela(page)      # volta para uma tela limpa
                except Exception:
                    pass

            if registros is None:
                motivo = erro if erro else "tabela do dia não apareceu"
                print(f"Dia {dia}: NÃO LIDO ({motivo})")
                falhas.append(dia)
                continue

            # Diagnóstico sem expor dados pessoais: mostra só os textos de status.
            if dia == DIA_INICIO and registros:
                print("AMOSTRA de status:", [r["status"][:60] for r in registros[:3]])

            # conta quantas vezes cada paciente aparece como atendido no dia
            contagem = {}   # nome normalizado -> [nome original, quantidade]
            for r in registros:
                if eh_atendido(r["status"]):
                    chave = normalizar(r["nome"])
                    if chave in contagem:
                        contagem[chave][1] += 1
                    else:
                        contagem[chave] = [r["nome"], 1]

            for nome, qtd in contagem.values():
                lin = achar_linha(ws, nome)
                ws.cell(row=lin, column=COL_DIA_1 + dia - 1, value="X" * qtd)

            repetidos = [n for n, q in contagem.values() if q > 1]
            wb.save(ARQUIVO)
            print(f"Dia {dia}: {len(contagem)} pacientes atendidos "
                  f"({len(registros)} registros na tabela)"
                  + (f" | repetidos (XX): {len(repetidos)}" if repetidos else ""))

        browser.close()

    if falhas:
        print("Dias NÃO lidos (confira na mão se têm atendimento):", falhas)


if __name__ == "__main__":
    main()
